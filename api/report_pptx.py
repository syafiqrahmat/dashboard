"""Renders the Development progress report as a .pptx.

The deck mimics api/LKTN_Enhancement_Progress.pptx: 16:9 slides where each
slide is ONE full-bleed PNG (no live text runs -- that keeps the layout
identical everywhere and sidesteps font embedding entirely). Every slide is
drawn with Pillow at SCALE=2 for crispness, then python-pptx drops the image
onto a blank layout at the same geometry the original deck uses
(0,38100 EMU / 16256000x9067800 EMU).

Copy is Malay, like the original, and the palette is the trio the original
deck leans on: slate (structure/neutral), amber (in progress), green (done).
"""
import io
import os
import unicodedata
from datetime import date, datetime

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

# Design canvas (logical px) and supersampling factor. All coordinates in
# the drawing helpers below are logical; _px() maps them to device pixels.
W, H = 1376, 768
SCALE = 2

M = 60                      # side margin
BANNER_Y, BANNER_H = 648, 60

SLATE = (100, 116, 139)
SLATE_DARK = (71, 85, 105)
AMBER = (245, 185, 66)
GREEN = (47, 163, 107)
INK = (30, 41, 59)
MUTED = (148, 163, 184)
BG = (244, 246, 247)
CARD = (255, 255, 255)
TRACK = (226, 232, 240)
GRID = (233, 237, 242)
WHITE = (255, 255, 255)

GREEN_BG = (226, 245, 236)
AMBER_BG = (254, 243, 220)
SLATE_BG = (237, 240, 244)

MONTHS_MS = ["Jan", "Feb", "Mac", "Apr", "Mei", "Jun",
             "Jul", "Ogo", "Sep", "Okt", "Nov", "Dis"]

STATUS_MS = {
    "completed": "Selesai",
    "inprogress": "Sedang Berjalan",
    "notstarted": "Belum Bermula",
}

_FONT_CACHE = {}


def _px(v):
    return int(round(v * SCALE))


def _font_files(bold):
    """Candidate font files, Windows first, then matplotlib's bundled
    DejaVu so the report still renders on a machine without Office fonts."""
    names = ["segoeuib.ttf", "tahomabd.ttf", "arialbd.ttf"] if bold else \
            ["segoeui.ttf", "tahoma.ttf", "arial.ttf"]
    out = [os.path.join("C:\\Windows\\Fonts", n) for n in names]
    try:
        import matplotlib
        ttf = os.path.join(os.path.dirname(matplotlib.__file__),
                           "mpl-data", "fonts", "ttf")
        out.append(os.path.join(ttf, "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"))
    except Exception:
        pass
    return out


def _font(size, bold=False):
    key = (round(size, 1), bold)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    font = None
    for path in _font_files(bold):
        if os.path.exists(path):
            font = ImageFont.truetype(path, _px(size))
            break
    if font is None:
        font = ImageFont.load_default()
    _FONT_CACHE[key] = font
    return font


def _status_key(pct, status_text=None):
    if status_text:
        norm = str(status_text).lower().replace(" ", "")
        if norm in STATUS_MS:
            return norm
    if pct is None:
        return "notstarted"
    if pct >= 100:
        return "completed"
    if pct <= 0:
        return "notstarted"
    return "inprogress"


def _status_color(key):
    return {"completed": GREEN, "inprogress": AMBER}.get(key, SLATE)


def _status_bg(key):
    return {"completed": GREEN_BG, "inprogress": AMBER_BG}.get(key, SLATE_BG)


def _fmt_pct(pct):
    if pct is None:
        return "-"
    if abs(pct - round(pct)) < 0.05:
        return f"{int(round(pct))}%"
    return f"{pct:.1f}%"


def _dt(text):
    """dd/mm/yyyy -> date, or None."""
    if not text:
        return None
    try:
        return datetime.strptime(str(text).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


class Canvas:
    def __init__(self, bg=BG, grid=False):
        self.img = Image.new("RGB", (_px(W), _px(H)), bg)
        self.d = ImageDraw.Draw(self.img, "RGBA")
        if grid:
            step = 64
            for x in range(step, W, step):
                self.d.line([_px(x), 0, _px(x), _px(H)], fill=GRID + (110,), width=_px(1))
            for y in range(step, H, step):
                self.d.line([0, _px(y), _px(W), _px(y)], fill=GRID + (110,), width=_px(1))

    def rrect(self, xy, r, fill, outline=None, width=1, shadow=False):
        x0, y0, x1, y1 = xy
        if x1 <= x0 or y1 <= y0:
            return
        r = max(1, min(r, (x1 - x0) / 2, (y1 - y0) / 2))
        if shadow:
            off = 5
            self.d.rounded_rectangle(
                [_px(x0) + _px(1), _px(y0) + _px(off), _px(x1) + _px(1), _px(y1) + _px(off)],
                radius=_px(r), fill=(15, 23, 42, 22))
        self.d.rounded_rectangle(
            [_px(x0), _px(y0), _px(x1), _px(y1)], radius=_px(r),
            fill=fill, outline=outline, width=_px(width))

    def text(self, xy, s, font, fill, anchor="la"):
        self.d.text((_px(xy[0]), _px(xy[1])), s, font=font, fill=fill, anchor=anchor)

    def textw(self, s, font):
        return self.d.textlength(s, font=font) / SCALE

    def line(self, xy, fill, width=1):
        self.d.line([_px(v) for v in xy], fill=fill, width=_px(width))

    def circle(self, cx, cy, r, fill=None, outline=None, width=1):
        self.d.ellipse([_px(cx - r), _px(cy - r), _px(cx + r), _px(cy + r)],
                       fill=fill, outline=outline, width=_px(width))


def _truncate(c, s, font, max_w):
    if c.textw(s, font) <= max_w:
        return s
    ell = "…"
    while s and c.textw(s + ell, font) > max_w:
        s = s[:-1]
    return s.rstrip() + ell


def _wrap(c, s, font, max_w):
    lines, cur = [], ""
    for word in str(s).split():
        trial = word if not cur else cur + " " + word
        if c.textw(trial, font) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines or [""]


def _wrap_segments(c, segments, reg, bold, max_w):
    """Wrap [(text, is_bold)] into lines of [(word, is_bold)].

    Pure-punctuation words (".", ",") glue onto the previous word instead of
    standing alone -- otherwise the space-before-token shows up as
    "44.6% ." in the banner."""
    tokens = []
    for text, is_bold in segments:
        for word in str(text).split():
            if tokens and word and all(ch in ".,;:!?" for ch in word):
                tokens[-1] = (tokens[-1][0] + word, tokens[-1][1])
            else:
                tokens.append((word, is_bold))
    space = c.textw(" ", reg)
    lines, cur, cur_w = [], [], 0.0
    for word, is_bold in tokens:
        f = bold if is_bold else reg
        w = c.textw(word, f)
        add = w if not cur else w + space
        if not cur or cur_w + add <= max_w:
            cur.append((word, is_bold))
            cur_w += add
        else:
            lines.append(cur)
            cur, cur_w = [(word, is_bold)], w
    if cur:
        lines.append(cur)
    return lines


def _draw_segments_line(c, line, cx, y, reg, bold, fill, fill_bold):
    space = c.textw(" ", reg)
    widths = [c.textw(w, bold if b else reg) for w, b in line]
    total = sum(widths) + space * (len(line) - 1)
    x = cx - total / 2
    for (word, is_bold), w in zip(line, widths):
        f = bold if is_bold else reg
        c.text((x, y), word, f, fill_bold if is_bold else fill, anchor="lm")
        x += w + space


def _donut(c, cx, cy, r, pct, color, width=22, track=TRACK):
    box = [_px(cx - r), _px(cy - r), _px(cx + r), _px(cy + r)]
    c.d.arc(box, 0, 360, fill=track, width=_px(width))
    pct = max(0.0, min(float(pct or 0), 100.0))
    if pct > 0:
        c.d.arc(box, -90, -90 + 360 * pct / 100.0, fill=color, width=_px(width))


def _pill(c, x, y, text, key, size=14, h=28, align="left"):
    font = _font(size, True)
    padx = 14
    w = c.textw(text, font) + padx * 2
    if align == "right":
        x = x - w
    c.rrect([x, y, x + w, y + h], h / 2, _status_bg(key))
    c.text((x + w / 2, y + h / 2), text, font, _status_color(key), anchor="mm")
    return w


def _draw_title(c, title, sub=None, stamp=None):
    c.text((W / 2, 44), title, _font(32, True), INK, anchor="ma")
    if sub:
        c.text((W / 2, 92), sub, _font(15), MUTED, anchor="ma")
    if stamp:
        c.text((W - M, 46), stamp, _font(12), MUTED, anchor="ra")


def _banner(c, segments):
    """Bottom narrative banner (slate, bold highlights), like the original."""
    c.rrect([M, BANNER_Y, W - M, BANNER_Y + BANNER_H], 16, SLATE, shadow=True)
    reg, bold = _font(15), _font(15, True)
    lines = _wrap_segments(c, segments, reg, bold, W - 2 * M - 80)
    lh = 22
    top = BANNER_Y + BANNER_H / 2 - lh * len(lines) / 2
    for i, line in enumerate(lines):
        _draw_segments_line(c, line, W / 2, top + lh * i + lh / 2, reg, bold,
                            (219, 229, 238), WHITE)


# ---------------------------------------------------------------- slides ---

def _draw_cover(c, report):
    # Faint tech-illustration stand-in: concentric rings, bottom right.
    for r, col in ((230, SLATE), (176, AMBER), (126, GREEN), (80, SLATE)):
        c.circle(W - 150, 430, r, outline=col + (46,), width=10)
    c.circle(W - 150, 430, 44, outline=GREEN + (70,), width=6)

    c.text((M, 200), "Laporan Kemajuan Sistem:", _font(46, True), INK, anchor="la")
    c.text((M, 264), "Pembangunan (Development)", _font(46, True), SLATE, anchor="la")

    only = (report.get("clients") or [None])[0] if len(report.get("clients") or []) == 1 else None
    card_b = 516 if only else 486
    c.rrect([M, 372, M + 660, card_b], 16, CARD, shadow=True)
    c.rrect([M, 372, M + 10, card_b], 5, GREEN)
    c.text((M + 34, 402), f"Dijana pada {report.get('generated_at') or '-'}",
           _font(20, True), INK, anchor="la")
    s = report.get("summary") or {}
    c.text((M + 34, 440),
           f"{s.get('total', 0)} klien · {report.get('module_count', 0)} modul · "
           f"{report.get('milestone_count', 0)} milestone",
           _font(16), SLATE, anchor="la")
    if only:
        # Per-client deck: name the client on the cover.
        c.text((M + 34, 474),
               _truncate(c, f"Klien: {only.get('client') or '-'} · "
                            f"{only.get('projek_name') or '-'} · ID {only.get('projek_id') or '-'}",
                         _font(15), 660 - 68),
               _font(15), SLATE, anchor="la")

    c.text((M, H - 60), "Dijana automatik daripada dashboard — jangan sunting tangan.",
           _font(13), MUTED, anchor="la")


def _draw_gauges(c, report):
    clients = report.get("clients") or []
    _draw_title(c, "Kemajuan Klien Development",
                "Purata Percentage daripada baris Project Details setiap klien",
                stamp=report.get("generated_at"))

    n = max(len(clients), 1)
    rows = 1 if n <= 3 else 2
    per_row = (n + rows - 1) // rows
    top, bottom = 128, 632
    gap_x, gap_y = 30, 24
    row_h = (bottom - top - gap_y * (rows - 1)) / rows

    for i, cl in enumerate(clients):
        r, col = divmod(i, per_row)
        count = min(per_row, n - r * per_row)
        # Cap the card width and centre the row: with one client a
        # full-bleed card looks empty, a ~560px card reads as a stat panel.
        cw = min((W - 2 * M - gap_x * (count - 1)) / count, 560)
        row_w = count * cw + gap_x * (count - 1)
        x0 = M + ((W - 2 * M) - row_w) / 2 + col * (cw + gap_x)
        y0 = top + r * (row_h + gap_y)
        x1, y1 = x0 + cw, y0 + row_h

        key = _status_key(cl.get("overall_pct"), cl.get("progress_status"))
        c.rrect([x0, y0, x1, y1], 18, CARD, shadow=True)

        c.text(((x0 + x1) / 2, y0 + 34), _truncate(c, cl.get("client") or "-",
               _font(23, True), cw - 40), _font(23, True), INK, anchor="ma")
        c.text(((x0 + x1) / 2, y0 + 64),
               _truncate(c, cl.get("projek_name") or "-", _font(14), cw - 40),
               _font(14), MUTED, anchor="ma")

        pct = cl.get("overall_pct")
        cx, cy = (x0 + x1) / 2, y0 + row_h / 2 + 8
        _donut(c, cx, cy, 74, pct, _status_color(key), width=24)
        c.text((cx, cy - 8), _fmt_pct(pct), _font(34, True), INK, anchor="mm")
        c.text((cx, cy + 24), "kemajuan", _font(13), MUTED, anchor="mm")

        _pill(c, (x0 + x1) / 2, y1 - 74, STATUS_MS[key], key, align="center")
        if cl.get("technology"):
            c.text((cx, y1 - 36), cl.get("technology"), _font(12), MUTED, anchor="ma")

    s = report.get("summary") or {}
    _banner(c, [
        ("Daripada ", False), (f"{s.get('total', 0)} klien", True),
        (" Development: ", False), (f"{s.get('completed', 0)} Selesai", True),
        (", ", False), (f"{s.get('in_progress', 0)} Sedang Berjalan", True),
        (", ", False), (f"{s.get('not_started', 0)} Belum Bermula", True),
        (". Purata kemajuan keseluruhan ", False),
        (f"{report.get('avg_pct') or '-'}%", True), (".", False),
    ])


def _draw_timeline(c, report):
    clients = report.get("clients") or []
    _draw_title(c, "Garis Masa Pembangunan",
                "Tempoh kontrak setiap klien Development", stamp=report.get("generated_at"))

    spans = []
    for cl in clients:
        s, e = _dt(cl.get("start_date")), _dt(cl.get("end_date"))
        if s and e:
            spans.append((s, e))
    if not spans:
        c.text((W / 2, 340), "Tiada tarikh mula/tamat direkod.", _font(18), MUTED, anchor="ma")
        _banner(c, [("Tiada data garis masa untuk dipaparkan.", True)])
        return

    from datetime import timedelta
    min_d = min(s for s, _ in spans) - timedelta(days=20)
    max_d = max(e for _, e in spans) + timedelta(days=20)
    span = max((max_d - min_d).days, 1)
    x0, x1 = 330, W - M - 70
    plot_w = x1 - x0

    def x_of(d):
        return x0 + (d - min_d).days / span * plot_w

    # Month grid + labels
    axis_y, grid_top, grid_bottom = 138, 166, 556
    months = (max_d.year - min_d.year) * 12 + (max_d.month - min_d.month) + 1
    step = 1
    while months / step * 66 > plot_w and step < 24:
        step += 1
    y_, m_ = min_d.year, min_d.month
    while (y_, m_) <= (max_d.year, max_d.month):
        tick = date(y_, m_, 1)
        if min_d <= tick <= max_d:
            x = x_of(tick)
            c.line([x, grid_top, x, grid_bottom], GRID, 1)
            if ((y_ - min_d.year) * 12 + m_ - min_d.month) % step == 0:
                c.text((x, axis_y), f"{MONTHS_MS[m_ - 1]} {y_}",
                       _font(12, True), SLATE, anchor="ma")
        m_ += 1
        if m_ > 12:
            m_, y_ = 1, y_ + 1

    # Today marker -- drawn before the bars so date labels stay on top
    today = date.today()
    if min_d <= today <= max_d:
        tx = x_of(today)
        y = grid_top
        while y < grid_bottom:
            c.line([tx, y, tx, min(y + 9, grid_bottom)], SLATE_DARK, 2)
            y += 15
        c.rrect([tx - 44, 112, tx + 44, 134], 11, SLATE_DARK)
        c.text((tx, 123), "Hari Ini", _font(12, True), WHITE, anchor="mm")

    # Contract bars -- cap the row height and centre the stack so a
    # single-client deck doesn't stretch one bar over the whole grid.
    n = len(clients)
    row_h = min((grid_bottom - grid_top) / max(n, 1), 120)
    y_start = grid_top + ((grid_bottom - grid_top) - row_h * n) / 2
    for i, cl in enumerate(clients):
        s, e = _dt(cl.get("start_date")), _dt(cl.get("end_date"))
        if not (s and e):
            continue
        cy = y_start + row_h * i + row_h / 2
        key = _status_key(cl.get("overall_pct"), cl.get("progress_status"))
        c.text((M, cy - 11), _truncate(c, cl.get("client") or "-", _font(18, True), 250),
               _font(18, True), INK, anchor="la")
        c.text((M, cy + 14), _truncate(c, cl.get("projek_name") or "-", _font(13), 250),
               _font(13), MUTED, anchor="la")

        bx0, bx1 = x_of(s), x_of(e)
        c.rrect([bx0, cy - 15, bx1, cy + 15], 15, TRACK)
        pct = cl.get("overall_pct") or 0
        fill_w = (bx1 - bx0) * max(0.0, min(float(pct), 100.0)) / 100.0
        if fill_w > 2:
            c.rrect([bx0, cy - 15, bx0 + max(fill_w, 30), cy + 15], 15, _status_color(key))
        c.text((bx0, cy + 26), cl.get("start_date") or "", _font(11), MUTED, anchor="la")
        c.text((bx1, cy + 26), cl.get("end_date") or "", _font(11), MUTED, anchor="ra")
        c.text((bx1 + 10, cy), _fmt_pct(cl.get("overall_pct")),
               _font(14, True), _status_color(key), anchor="lm")

    # Narrative banner
    _banner(c, [
        ("Tempoh projek merentasi ", False),
        (f"{MONTHS_MS[min(s for s, _ in spans).month - 1]} {min(s for s, _ in spans).year} hingga "
         f"{MONTHS_MS[max(e for _, e in spans).month - 1]} {max(e for _, e in spans).year}", True),
        (". Penanda hari ini ", False),
        (today.strftime("%d/%m/%Y"), True),
        (" menunjukkan klien masih dalam fasa pembangunan.", False),
    ])


def _bar_row(c, x, y, w, label, pct, row_h, label_w):
    key = _status_key(pct)
    font = _font(13)
    c.text((x, y + row_h / 2), _truncate(c, label, font, label_w), font, INK, anchor="lm")
    tx0 = x + label_w + 12
    tx1 = x + w - 52
    c.rrect([tx0, y + row_h / 2 - 6, tx1, y + row_h / 2 + 6], 6, TRACK)
    p = max(0.0, min(float(pct or 0), 100.0))
    if p > 0:
        fw = (tx1 - tx0) * p / 100.0
        c.rrect([tx0, y + row_h / 2 - 6, tx0 + max(fw, 12), y + row_h / 2 + 6], 6,
                _status_color(key))
    c.text((x + w, y + row_h / 2), _fmt_pct(pct), _font(13, True),
           _status_color(key), anchor="rm")


def _draw_client(c, report, cl):
    key = _status_key(cl.get("overall_pct"), cl.get("progress_status"))
    _draw_title(c, f"Status Modul: {cl.get('projek_name') or '-'}",
                f"{cl.get('client')} · ID {cl.get('projek_id') or '-'} · "
                f"{cl.get('technology') or '-'}",
                stamp=report.get("generated_at"))

    # Profile card
    px0, py0, px1, py1 = M, 124, M + 380, 440
    c.rrect([px0, py0, px1, py1], 16, CARD, shadow=True)
    c.text((px0 + 26, py0 + 30), "Profil Projek", _font(17, True), INK, anchor="la")
    fields = [
        ("ID Projek", cl.get("projek_id")),
        ("Teknologi", cl.get("technology")),
        ("Mula", cl.get("start_date")),
        ("Tamat", cl.get("end_date")),
    ]
    fy = py0 + 64
    for label, val in fields:
        c.text((px0 + 26, fy), label, _font(13), MUTED, anchor="la")
        c.text((px1 - 26, fy), _truncate(c, val or "-", _font(14, True), 190),
               _font(14, True), INK, anchor="ra")
        c.line([px0 + 26, fy + 24, px1 - 26, fy + 24], TRACK, 1)
        fy += 44
    c.text((px0 + 26, fy + 4), "Status", _font(13), MUTED, anchor="la")
    _pill(c, px0 + 92, fy - 8, STATUS_MS[key], key)
    c.text((px0 + 26, fy + 34), _fmt_pct(cl.get("overall_pct")),
           _font(34, True), _status_color(key), anchor="la")
    c.text((px0 + 140, fy + 42), "kemajuan\nkeseluruhan", _font(12), MUTED, anchor="la")

    # Module bars card (project modules + task-detail modules side by side)
    bx0, by0, bx1, by1 = M + 410, 124, W - M, 440
    c.rrect([bx0, by0, bx1, by1], 16, CARD, shadow=True)
    modules = cl.get("modules") or []
    td = cl.get("task_detail_modules") or []
    avail = (by1 - by0) - 58 - 14

    if modules and td:
        col_w = (bx1 - bx0 - 78) / 2
        groups = [((bx0 + 24), col_w, "Project Details", modules),
                  ((bx0 + 54 + col_w), col_w, "Task Detail", td)]
    else:
        groups = [((bx0 + 24), bx1 - bx0 - 48, "Kemajuan Modul", modules or td)]

    if not modules and not td:
        c.text(((bx0 + bx1) / 2, (by0 + by1) / 2),
               "Tiada baris modul untuk projek ini.", _font(16), MUTED, anchor="mm")

    for gx, gw, header, items in groups:
        c.text((gx, by0 + 30), header, _font(16, True), INK, anchor="la")
        c.line([gx, by0 + 48, gx + gw, by0 + 48], TRACK, 1)
        if not items:
            c.text((gx, by0 + 80), "Tiada data.", _font(13), MUTED, anchor="la")
            continue
        cap = int(avail // 26)
        shown, overflow = (items[:cap], len(items) - cap) if len(items) > cap else (items, 0)
        row_h = min(30, avail / max(len(shown) + (0.45 if overflow else 0), 1))
        ry = by0 + 60
        label_w = max(90, gw * 0.34)
        for it in shown:
            _bar_row(c, gx, ry, gw, it.get("name") or "-", it.get("pct"), row_h, label_w)
            ry += row_h
        if overflow:
            c.text((gx, ry + 8), f"+ {overflow} lagi…", _font(12, True), MUTED, anchor="la")

    # Milestones as chips
    ms = cl.get("milestones") or []
    mx0, my0, mx1, my1 = M, 458, W - M, 636
    c.rrect([mx0, my0, mx1, my1], 16, CARD, shadow=True)
    c.text((mx0 + 24, my0 + 26), f"Pencapaian (Milestone) — {len(ms)}",
           _font(16, True), INK, anchor="la")
    if not ms:
        c.text((mx0 + 24, my0 + 64), "Tiada milestone direkod untuk projek ini.",
               _font(13), MUTED, anchor="la")
    else:
        chip_w, chip_h, gap = 234, 32, 8
        per_row = int((mx1 - mx0 - 48 + gap) // (chip_w + gap))
        cap = per_row * 3
        shown, overflow = (ms[:cap], len(ms) - cap) if len(ms) > cap else (ms, 0)
        cy = my0 + 54
        for i, m in enumerate(shown):
            r, col = divmod(i, per_row)
            if r > 2:
                break
            x = mx0 + 24 + col * (chip_w + gap)
            y = cy + r * (chip_h + gap)
            mk = _status_key(m.get("progress"))
            c.rrect([x, y, x + chip_w, y + chip_h], 9, _status_bg(mk))
            c.circle(x + 17, y + chip_h / 2, 7, fill=_status_color(mk))
            if mk == "completed":
                c.line([x + 13, y + chip_h / 2 + 1, x + 16, y + chip_h / 2 + 4], WHITE, 2)
                c.line([x + 16, y + chip_h / 2 + 4, x + 21, y + chip_h / 2 - 3], WHITE, 2)
            else:
                c.line([x + 13, y + chip_h / 2, x + 21, y + chip_h / 2], WHITE, 2)
            name = _truncate(c, m.get("task_name") or "-", _font(13, True), chip_w - 92)
            c.text((x + 32, y + chip_h / 2), name, _font(13, True), INK, anchor="lm")
            c.text((x + chip_w - 12, y + chip_h / 2), _fmt_pct(m.get("progress")),
                   _font(13, True), _status_color(mk), anchor="rm")
        if overflow:
            c.text((mx1 - 24, my0 + 26), f"+ {overflow} lagi…",
                   _font(12, True), MUTED, anchor="ra")

    # Narrative banner
    mods = [m for m in (modules or td) if m.get("pct") is not None]
    segments = [(f"{cl.get('client')}: ", True)]
    if mods:
        hi = max(mods, key=lambda m: m["pct"])
        lo = min(mods, key=lambda m: m["pct"])
        segments += [
            (f"{len(mods)} modul, kemajuan ", False),
            (_fmt_pct(cl.get("overall_pct")), True),
            (". Tertinggi ", False), (f"{hi['name']} ({_fmt_pct(hi['pct'])})", True),
            (", terendah ", False), (f"{lo['name']} ({_fmt_pct(lo['pct'])})", True),
            (".", False),
        ]
    else:
        segments += [("kemajuan keseluruhan ", False),
                     (_fmt_pct(cl.get("overall_pct")), True),
                     (" pada status ", False), (STATUS_MS[key], True), (".", False)]
    _banner(c, segments)


def _draw_summary(c, report):
    clients = report.get("clients") or []
    s = report.get("summary") or {}
    _draw_title(c, "Rumusan & Langkah Seterusnya",
                "Ringkasan status klien Development", stamp=report.get("generated_at"))

    buckets = [("completed", "Selesai", GREEN),
               ("in_progress", "Sedang Berjalan", AMBER),
               ("not_started", "Belum Bermula", SLATE)]
    # summary counts use in_progress/not_started; _status_key() returns
    # inprogress/notstarted -- map before matching members to a bucket.
    bkey_to_status = {"completed": "completed", "in_progress": "inprogress",
                      "not_started": "notstarted"}
    gap = 32
    cw = (W - 2 * M - gap * 2) / 3
    for i, (bkey, label, col) in enumerate(buckets):
        x0 = M + i * (cw + gap)
        y0, y1 = 134, 372
        c.rrect([x0, y0, x0 + cw, y1], 16, CARD, shadow=True)
        c.rrect([x0, y0, x0 + cw, y0 + 66], 16, col)
        c.d.rectangle([_px(x0), _px(y0 + 40), _px(x0 + cw), _px(y0 + 66)], fill=col)
        c.text((x0 + cw / 2, y0 + 33), label, _font(19, True), WHITE, anchor="mm")
        c.text((x0 + cw / 2, y0 + 118), str(s.get(bkey, 0)),
               _font(54, True), col, anchor="mm")
        members = [cl.get("client") or "-" for cl in clients
                   if _status_key(cl.get("overall_pct"), cl.get("progress_status"))
                   == bkey_to_status[bkey]]
        text = " · ".join(members) if members else "—"
        lines = _wrap(c, text, _font(14), cw - 48)
        ly = y0 + 176
        for line in lines[:4]:
            c.text((x0 + cw / 2, ly), line, _font(14), SLATE, anchor="ma")
            ly += 22

    # Next steps
    nx0, ny0, nx1, ny1 = M, 396, W - M, 632
    c.rrect([nx0, ny0, nx1, ny1], 16, CARD, shadow=True)
    c.text((nx0 + 26, ny0 + 32), "Langkah Seterusnya", _font(18, True), INK, anchor="la")
    steps = []
    for cl in clients:
        mods = [m for m in (cl.get("modules") or []) if m.get("pct") is not None]
        if not mods:
            continue
        low = min(mods, key=lambda m: m["pct"])
        if (low.get("pct") or 0) < 100:
            steps.append(
                (cl,
                 f"Segerakan modul {low['name']} ({_fmt_pct(low['pct'])}) untuk "
                 f"{cl.get('client')} — tamat {cl.get('end_date') or '-'}"))
    if not steps:
        steps = [(None, "Semua modul telah mencapai 100% — sedia untuk serahan UAT/FAT.")]
    sy = ny0 + 74
    for cl, text in steps[:4]:
        key = _status_key(cl.get("overall_pct") if cl else None,
                          cl.get("progress_status") if cl else "completed")
        c.circle(nx0 + 34, sy + 10, 6, fill=_status_color(key))
        for j, line in enumerate(_wrap(c, text, _font(15), nx1 - nx0 - 90)[:2]):
            c.text((nx0 + 54, sy + 10 + j * 22), line, _font(15), INK, anchor="la")
            if j == 0:
                sy += 0
        sy += 46
    if len(steps) > 4:
        c.text((nx1 - 26, ny0 + 32), f"+ {len(steps) - 4} lagi…",
               _font(12, True), MUTED, anchor="ra")

    _banner(c, [
        ("Purata kemajuan keseluruhan ", False),
        (f"{report.get('avg_pct') or '-'}%", True), (" merentasi ", False),
        (f"{s.get('total', 0)} klien", True), (" Development: ", False),
        (f"{s.get('completed', 0)} Selesai", True), (", ", False),
        (f"{s.get('in_progress', 0)} Sedang Berjalan", True), (", ", False),
        (f"{s.get('not_started', 0)} Belum Bermula", True), (".", False),
    ])


def render_slide(kind, data):
    """Render one slide and return a PIL image (1376x768 * SCALE)."""
    if kind == "cover":
        c = Canvas(grid=False)
        _draw_cover(c, data)
    elif kind == "gauges":
        c = Canvas(grid=True)
        _draw_gauges(c, data)
    elif kind == "timeline":
        c = Canvas(grid=True)
        _draw_timeline(c, data)
    elif kind == "client":
        c = Canvas(grid=True)
        _draw_client(c, data["report"], data["client"])
    elif kind == "summary":
        c = Canvas(grid=True)
        _draw_summary(c, data)
    else:
        raise ValueError(f"unknown slide kind: {kind}")
    return c.img


def _png_bytes(img):
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


def build_development_deck(report):
    """report (db.fetch_development_report_data) -> BytesIO holding the .pptx."""
    report = dict(report or {})
    report["module_count"] = sum(len(c.get("modules") or []) for c in report.get("clients") or [])
    report["milestone_count"] = sum(len(c.get("milestones") or []) for c in report.get("clients") or [])
    pcts = [c.get("overall_pct") for c in report.get("clients") or []
            if c.get("overall_pct") is not None]
    report["avg_pct"] = round(sum(pcts) / len(pcts), 1) if pcts else None

    slides = [("cover", report), ("gauges", report), ("timeline", report)]
    slides += [("client", {"report": report, "client": cl})
               for cl in report.get("clients") or []]
    slides.append(("summary", report))

    prs = Presentation()
    prs.slide_width = Emu(16256000)
    prs.slide_height = Emu(9144000)
    blank = prs.slide_layouts[6]
    for kind, data in slides:
        img = render_slide(kind, data)
        slide = prs.slides.add_slide(blank)
        # Same geometry as the original deck: full-bleed image inset 38100 EMU.
        slide.shapes.add_picture(_png_bytes(img), Emu(0), Emu(38100),
                                 width=Emu(16256000), height=Emu(9067800))

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out
